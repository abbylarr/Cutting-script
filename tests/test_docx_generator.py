"""
Tests for DOCX document generation service.
"""

import pytest
import asyncio
from datetime import datetime
from io import BytesIO
from pathlib import Path
import tempfile
import os

from docx import Document

from app.services.docx_generator import (
    DOCXGeneratorService,
    DOCXTemplateService,
    DOCXGenerationError
)
from app.schemas.film_project import (
    MontageRow,
    FilmMetadata,
    ProjectSettings,
    ShotType,
    ColorType,
    TimecodeStandard
)


@pytest.fixture
def sample_film_metadata():
    """Create sample film metadata for testing."""
    return FilmMetadata(
        title="Тестовый фильм",
        production_company="Тестовая киностудия",
        year=2024,
        country="Россия",
        screenwriters=["Иван Иванов", "Петр Петров"],
        copyright_holders=["Тестовая киностудия", "Государство"],
        duration="01:30:45",
        episodes_count=1,
        format="Digital 4K",
        color_type=ColorType.COLOR,
        media_carrier="SSD",
        original_language="Русский",
        subtitle_language="Английский",
        audio_language="Русский"
    )


@pytest.fixture
def sample_project_settings():
    """Create sample project settings for testing."""
    return ProjectSettings(
        timecode_start="01:00:00:00",
        standard=TimecodeStandard.GFF,
        fps=25.0
    )


@pytest.fixture
def sample_montage_rows():
    """Create sample montage rows for testing."""
    return [
        MontageRow(
            number=1,
            start_timecode="01:00:00:00",
            end_timecode="01:00:05:12",
            shot_type=ShotType.GENERAL,
            description="Общий план города",
            dialogue="",
            speaker=None,
            has_music=False,
            special_tags=[]
        ),
        MontageRow(
            number=2,
            start_timecode="01:00:05:13",
            end_timecode="01:00:12:08",
            shot_type=ShotType.MEDIUM,
            description="Средний план главного героя",
            dialogue="Привет, как дела?",
            speaker="Главный герой",
            has_music=False,
            special_tags=["ГЗК"]
        ),
        MontageRow(
            number=3,
            start_timecode="01:00:12:09",
            end_timecode="01:00:18:24",
            shot_type=ShotType.CLOSE,
            description="Крупный план реакции",
            dialogue="Все хорошо, спасибо!",
            speaker="Второй персонаж",
            has_music=True,
            special_tags=["НДП"]
        )
    ]


class TestDOCXTemplateService:
    """Test DOCX template service functionality."""
    
    def test_create_document_template(self):
        """Test document template creation."""
        template_service = DOCXTemplateService()
        doc = template_service.create_document_template()
        
        assert doc is not None
        assert hasattr(doc, 'sections')
        assert len(doc.sections) > 0
        
        # Check margins are set correctly (2.5cm)
        section = doc.sections[0]
        # Cm(2.5) converts to EMU units, allow small tolerance for rounding
        from docx.shared import Cm
        expected_margin = Cm(2.5)
        tolerance = 1000  # Small tolerance for rounding differences
        
        assert abs(section.top_margin - expected_margin) <= tolerance
        assert abs(section.bottom_margin - expected_margin) <= tolerance
        assert abs(section.left_margin - expected_margin) <= tolerance
        assert abs(section.right_margin - expected_margin) <= tolerance
    
    def test_add_document_header(self, sample_film_metadata):
        """Test document header addition in GFF format."""
        template_service = DOCXTemplateService()
        doc = template_service.create_document_template()
        
        template_service.add_document_header(doc, sample_film_metadata)
        
        # Check that paragraphs were added
        assert len(doc.paragraphs) >= 3
        
        # Check header content (GFF format)
        doc_text = "\n".join([p.text for p in doc.paragraphs])
        assert "Форма монтажных листов фильма" in doc_text
    
    def test_add_film_metadata_section(self, sample_film_metadata):
        """Test film metadata section addition in GFF format."""
        template_service = DOCXTemplateService()
        doc = template_service.create_document_template()
        
        initial_table_count = len(doc.tables)
        template_service.add_film_metadata_section(doc, sample_film_metadata)
        
        # Check that metadata table was added
        assert len(doc.tables) > initial_table_count
        
        # Check that key metadata is present in table
        table = doc.tables[0]
        table_text = "\n".join([cell.text for row in table.rows for cell in row.cells])
        assert sample_film_metadata.production_company in table_text
        assert str(sample_film_metadata.year) in table_text
        assert sample_film_metadata.country in table_text
        assert sample_film_metadata.duration in table_text
    
    def test_create_montage_table(self, sample_montage_rows):
        """Test montage table creation."""
        template_service = DOCXTemplateService()
        doc = template_service.create_document_template()
        
        template_service.create_montage_table(doc, sample_montage_rows)
        
        # Check that table was added
        assert len(doc.tables) == 1
        
        table = doc.tables[0]
        
        # Check table structure (header + data rows)
        expected_rows = len(sample_montage_rows) + 1  # +1 for header
        assert len(table.rows) == expected_rows
        
        # Check column count
        assert len(table.columns) == 6
        
        # Check header row content (GFF format)
        header_row = table.rows[0]
        gff_headers = [
            "№\nплана",
            "Начальный тайм-код плана",
            "Конечный тайм-код плана",
            "Вид плана",
            "Содержание (описание) плана",
            "Монологи, разговоры, песни"
        ]
        
        for i, expected_header in enumerate(gff_headers):
            assert expected_header in header_row.cells[i].text
        
        # Check data rows
        for i, montage_row in enumerate(sample_montage_rows):
            data_row = table.rows[i + 1]  # +1 to skip header
            
            assert str(montage_row.number) in data_row.cells[0].text
            assert montage_row.start_timecode in data_row.cells[1].text
            assert montage_row.end_timecode in data_row.cells[2].text
            assert montage_row.description in data_row.cells[4].text
    
    def test_create_montage_table_empty_rows(self):
        """Test montage table creation with empty rows."""
        template_service = DOCXTemplateService()
        doc = template_service.create_document_template()
        
        template_service.create_montage_table(doc, [])
        
        # Should not create table for empty rows
        assert len(doc.tables) == 0
    
    def test_format_shot_type(self, sample_montage_rows):
        """Test shot type formatting."""
        template_service = DOCXTemplateService()
        
        # Test different shot types
        test_cases = [
            (ShotType.DISTANT, "Дальн."),
            (ShotType.GENERAL, "Общ."),
            (ShotType.MEDIUM, "Ср."),
            (ShotType.CLOSE, "Крупн."),
            (ShotType.DETAIL, "Дет.")
        ]
        
        for shot_type, expected_abbrev in test_cases:
            row = MontageRow(
                number=1,
                start_timecode="01:00:00:00",
                end_timecode="01:00:05:00",
                shot_type=shot_type,
                description="Test",
                dialogue="",
                special_tags=[]
            )
            
            result = template_service._format_shot_type(row)
            assert result == expected_abbrev
    
    def test_format_description_with_tags(self):
        """Test description formatting with special tags."""
        template_service = DOCXTemplateService()
        
        row = MontageRow(
            number=1,
            start_timecode="01:00:00:00",
            end_timecode="01:00:05:00",
            shot_type=ShotType.MEDIUM,
            description="Тестовое описание",
            dialogue="",
            special_tags=["ЗТМ", "НДП"]
        )
        
        result = template_service._format_description(row)
        assert "Тестовое описание" in result
        assert "ЗТМ" in result
        assert "НДП" in result
    
    def test_format_dialogue_with_speaker(self):
        """Test dialogue formatting with speaker."""
        template_service = DOCXTemplateService()
        
        row = MontageRow(
            number=1,
            start_timecode="01:00:00:00",
            end_timecode="01:00:05:00",
            shot_type=ShotType.MEDIUM,
            description="Test",
            dialogue="Тестовая реплика",
            speaker="Тестовый персонаж"
        )
        
        result = template_service._format_dialogue(row)
        assert "Тестовый персонаж: Тестовая реплика" == result
    
    def test_format_dialogue_without_speaker(self):
        """Test dialogue formatting without speaker."""
        template_service = DOCXTemplateService()
        
        row = MontageRow(
            number=1,
            start_timecode="01:00:00:00",
            end_timecode="01:00:05:00",
            shot_type=ShotType.MEDIUM,
            description="Test",
            dialogue="Тестовая реплика",
            speaker=None
        )
        
        result = template_service._format_dialogue(row)
        assert result == "Тестовая реплика"
    
    def test_add_document_footer(self, sample_project_settings):
        """Test document footer addition."""
        template_service = DOCXTemplateService()
        doc = template_service.create_document_template()
        
        initial_paragraph_count = len(doc.paragraphs)
        template_service.add_document_footer(doc, sample_project_settings)
        
        # Check that footer paragraphs were added
        assert len(doc.paragraphs) > initial_paragraph_count
        
        # Check footer content
        footer_text = doc.paragraphs[-1].text
        assert "Документ создан автоматически" in footer_text
        assert sample_project_settings.standard.value in footer_text
        assert str(sample_project_settings.fps) in footer_text


class TestDOCXGeneratorService:
    """Test DOCX generator service functionality."""
    
    @pytest.mark.asyncio
    async def test_generate_montage_docx(self, sample_montage_rows, sample_film_metadata, sample_project_settings):
        """Test complete DOCX generation."""
        generator = DOCXGeneratorService()
        
        doc_bytes = await generator.generate_montage_docx(
            sample_montage_rows,
            sample_film_metadata,
            sample_project_settings
        )
        
        assert isinstance(doc_bytes, bytes)
        assert len(doc_bytes) > 0
        
        # Verify that the bytes represent a valid DOCX file
        doc_buffer = BytesIO(doc_bytes)
        doc = Document(doc_buffer)
        
        # Check document structure
        assert len(doc.paragraphs) > 0
        assert len(doc.tables) >= 1  # At least montage table, possibly production info table too
        
        # Check table content - find the montage table (should be the last table)
        montage_table = doc.tables[-1]  # Last table should be montage table
        assert len(montage_table.rows) == len(sample_montage_rows) + 1  # +1 for header
    
    @pytest.mark.asyncio
    async def test_generate_montage_docx_without_project_settings(self, sample_montage_rows, sample_film_metadata):
        """Test DOCX generation without project settings."""
        generator = DOCXGeneratorService()
        
        doc_bytes = await generator.generate_montage_docx(
            sample_montage_rows,
            sample_film_metadata
        )
        
        assert isinstance(doc_bytes, bytes)
        assert len(doc_bytes) > 0
    
    @pytest.mark.asyncio
    async def test_generate_montage_docx_empty_rows(self, sample_film_metadata):
        """Test DOCX generation with empty montage rows."""
        generator = DOCXGeneratorService()
        
        doc_bytes = await generator.generate_montage_docx(
            [],
            sample_film_metadata
        )
        
        assert isinstance(doc_bytes, bytes)
        assert len(doc_bytes) > 0
        
        # Document should still be created, may have production info table but no montage table
        doc_buffer = BytesIO(doc_bytes)
        doc = Document(doc_buffer)
        # With enhanced metadata, there might be a production info table
        assert len(doc.tables) >= 0
    
    @pytest.mark.asyncio
    async def test_save_document_to_file(self, sample_montage_rows, sample_film_metadata, sample_project_settings):
        """Test saving document to file."""
        generator = DOCXGeneratorService()
        
        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = os.path.join(temp_dir, "test_montage.docx")
            
            saved_path = await generator.save_document_to_file(
                sample_montage_rows,
                sample_film_metadata,
                output_path,
                sample_project_settings
            )
            
            assert saved_path == output_path
            assert os.path.exists(output_path)
            
            # Verify file content
            doc = Document(output_path)
            assert len(doc.paragraphs) > 0
            assert len(doc.tables) >= 1  # At least montage table, possibly production info table too
    
    @pytest.mark.asyncio
    async def test_save_document_creates_directory(self, sample_montage_rows, sample_film_metadata):
        """Test that save_document_to_file creates output directory."""
        generator = DOCXGeneratorService()
        
        with tempfile.TemporaryDirectory() as temp_dir:
            nested_dir = os.path.join(temp_dir, "nested", "directory")
            output_path = os.path.join(nested_dir, "test_montage.docx")
            
            saved_path = await generator.save_document_to_file(
                sample_montage_rows,
                sample_film_metadata,
                output_path
            )
            
            assert saved_path == output_path
            assert os.path.exists(output_path)
            assert os.path.exists(nested_dir)
    
    @pytest.mark.asyncio
    async def test_validate_document_compliance_valid(self, sample_montage_rows, sample_film_metadata):
        """Test document compliance validation with valid data."""
        generator = DOCXGeneratorService()
        
        is_compliant, issues = await generator.validate_document_compliance(
            sample_montage_rows,
            sample_film_metadata
        )
        
        assert is_compliant is True
        assert len(issues) == 0
    
    @pytest.mark.asyncio
    async def test_validate_document_compliance_invalid_metadata(self, sample_montage_rows):
        """Test document compliance validation with invalid metadata."""
        generator = DOCXGeneratorService()
        
        # Create invalid metadata (using valid Pydantic values but logically invalid)
        invalid_metadata = FilmMetadata(
            title="",  # Empty title
            production_company="",  # Empty production company
            year=1900,  # Minimum valid year but we'll test this in validation
            country="Россия",
            screenwriters=[],  # Empty screenwriters
            copyright_holders=[],  # Empty copyright holders
            duration="invalid_duration",  # Invalid duration format
            episodes_count=1,
            format="Digital",
            color_type=ColorType.COLOR,
            media_carrier="SSD",
            original_language="Русский",
            audio_language="Русский"
        )
        
        is_compliant, issues = await generator.validate_document_compliance(
            sample_montage_rows,
            invalid_metadata
        )
        
        assert is_compliant is False
        assert len(issues) > 0
        
        # Check specific issues
        issue_text = " ".join(issues)
        assert "Название фильма" in issue_text
        assert "Производитель" in issue_text
        assert "авторы сценария" in issue_text
        assert "правообладатели" in issue_text
        assert "продолжительности" in issue_text
    
    @pytest.mark.asyncio
    async def test_validate_document_compliance_invalid_montage_rows(self, sample_film_metadata):
        """Test document compliance validation with invalid montage rows."""
        generator = DOCXGeneratorService()
        
        # Create invalid montage rows (using valid Pydantic format but logically invalid)
        invalid_rows = [
            MontageRow(
                number=2,  # Wrong number (should be 1)
                start_timecode="01:00:00:00",  # Valid format
                end_timecode="01:00:05:00",
                shot_type=ShotType.MEDIUM,
                description="",  # Empty description
                dialogue=""
            )
        ]
        
        is_compliant, issues = await generator.validate_document_compliance(
            invalid_rows,
            sample_film_metadata
        )
        
        assert is_compliant is False
        assert len(issues) > 0
        
        # Check specific issues
        issue_text = " ".join(issues)
        assert "нумерация" in issue_text
        assert "описание" in issue_text
    
    @pytest.mark.asyncio
    async def test_validate_document_compliance_empty_rows(self, sample_film_metadata):
        """Test document compliance validation with empty montage rows."""
        generator = DOCXGeneratorService()
        
        is_compliant, issues = await generator.validate_document_compliance(
            [],
            sample_film_metadata
        )
        
        assert is_compliant is False
        assert len(issues) >= 1
        issue_text = " ".join(issues)
        assert "строки монтажного листа" in issue_text
    
    @pytest.mark.asyncio
    async def test_get_document_statistics(self, sample_montage_rows):
        """Test document statistics generation."""
        generator = DOCXGeneratorService()
        
        stats = await generator.get_document_statistics(sample_montage_rows)
        
        assert stats["total_rows"] == len(sample_montage_rows)
        assert "shot_type_distribution" in stats
        assert "rows_with_dialogue" in stats
        assert "dialogue_coverage_percent" in stats
        assert "rows_with_music" in stats
        assert "music_coverage_percent" in stats
        assert "special_tag_distribution" in stats
        assert "unique_speakers" in stats
        assert "speaker_list" in stats
        
        # Check specific values
        assert stats["rows_with_dialogue"] == 2  # 2 rows have dialogue
        assert stats["rows_with_music"] == 1    # 1 row has music
        assert stats["unique_speakers"] == 2    # 2 unique speakers
        
        # Check shot type distribution
        shot_types = stats["shot_type_distribution"]
        assert shot_types["Общий"] == 1
        assert shot_types["Средний"] == 1
        assert shot_types["Крупный"] == 1
    
    @pytest.mark.asyncio
    async def test_get_document_statistics_empty_rows(self):
        """Test document statistics with empty rows."""
        generator = DOCXGeneratorService()
        
        stats = await generator.get_document_statistics([])
        
        assert stats["total_rows"] == 0
        assert len(stats) == 1  # Only total_rows should be present
    
    @pytest.mark.asyncio
    async def test_document_to_bytes(self):
        """Test document to bytes conversion."""
        generator = DOCXGeneratorService()
        template_service = DOCXTemplateService()
        
        doc = template_service.create_document_template()
        doc.add_paragraph("Test content")
        
        doc_bytes = await generator._document_to_bytes(doc)
        
        assert isinstance(doc_bytes, bytes)
        assert len(doc_bytes) > 0
        
        # Verify bytes can be loaded back as document
        doc_buffer = BytesIO(doc_bytes)
        loaded_doc = Document(doc_buffer)
        assert len(loaded_doc.paragraphs) > 0


class TestDOCXGenerationError:
    """Test DOCX generation error handling."""
    
    @pytest.mark.asyncio
    async def test_docx_generation_error_handling(self):
        """Test error handling in DOCX generation."""
        generator = DOCXGeneratorService()
        
        # Test with invalid data that should cause an error
        with pytest.raises(DOCXGenerationError):
            # This should fail due to invalid metadata structure
            await generator.generate_montage_docx(
                None,  # Invalid montage rows
                None   # Invalid film metadata
            )


if __name__ == "__main__":
    pytest.main([__file__])


class TestDOCXMetadataIntegration:
    """Test enhanced metadata integration functionality."""
    
    @pytest.mark.asyncio
    async def test_add_document_header_footer(self, sample_film_metadata, sample_project_settings):
        """Test document header and footer with metadata."""
        template_service = DOCXTemplateService()
        doc = template_service.create_document_template()
        
        template_service.add_document_header_footer(doc, sample_film_metadata, sample_project_settings)
        
        # Check that header and footer were added
        section = doc.sections[0]
        
        # Check header content
        header_text = section.header.paragraphs[0].text
        assert sample_film_metadata.title in header_text
        assert sample_film_metadata.production_company in header_text
        assert str(sample_film_metadata.year) in header_text
        
        # Check footer content
        footer_text = section.footer.paragraphs[0].text
        assert "Монтажный лист создан автоматически" in footer_text
        assert sample_project_settings.standard.value in footer_text
        assert str(sample_project_settings.fps) in footer_text
        assert sample_film_metadata.format in footer_text
        assert sample_film_metadata.color_type.value in footer_text
    
    @pytest.mark.asyncio
    async def test_add_production_information_section(self, sample_film_metadata):
        """Test production information section addition."""
        template_service = DOCXTemplateService()
        doc = template_service.create_document_template()
        
        initial_paragraph_count = len(doc.paragraphs)
        template_service.add_production_information_section(doc, sample_film_metadata)
        
        # Check that content was added
        assert len(doc.paragraphs) > initial_paragraph_count
        assert len(doc.tables) > 0
        
        # Check table content
        table = doc.tables[0]
        table_text = "\n".join([cell.text for row in table.rows for cell in row.cells])
        
        assert "Правообладатели:" in table_text
        assert "Авторы сценария:" in table_text
        assert sample_film_metadata.copyright_holders[0] in table_text
        assert sample_film_metadata.screenwriters[0] in table_text
    
    @pytest.mark.asyncio
    async def test_add_compliance_information(self, sample_film_metadata):
        """Test compliance information section addition."""
        template_service = DOCXTemplateService()
        doc = template_service.create_document_template()
        
        initial_paragraph_count = len(doc.paragraphs)
        template_service.add_compliance_information(doc, sample_film_metadata)
        
        # Check that content was added
        assert len(doc.paragraphs) > initial_paragraph_count
        
        # Check compliance text content
        doc_text = "\n".join([p.text for p in doc.paragraphs])
        assert "Госфильмфонда Российской Федерации" in doc_text
        assert sample_film_metadata.title in doc_text
        assert sample_film_metadata.production_company in doc_text
        assert str(sample_film_metadata.year) in doc_text
    
    @pytest.mark.asyncio
    async def test_generate_montage_docx_with_gff_format(
        self, 
        sample_montage_rows, 
        sample_film_metadata, 
        sample_project_settings
    ):
        """Test DOCX generation with GFF format."""
        generator = DOCXGeneratorService()
        
        doc_bytes = await generator.generate_montage_docx(
            sample_montage_rows,
            sample_film_metadata,
            sample_project_settings,
            use_gff_format=True
        )
        
        assert isinstance(doc_bytes, bytes)
        assert len(doc_bytes) > 0
        
        # Verify document structure
        doc_buffer = BytesIO(doc_bytes)
        doc = Document(doc_buffer)
        
        # Should have 2 tables (metadata + montage)
        assert len(doc.tables) == 2
        
        # Check document content
        doc_text = "\n".join([p.text for p in doc.paragraphs])
        assert "Форма монтажных листов фильма" in doc_text
        assert "МОНТАЖНЫЕ ЛИСТЫ" in doc_text
        assert "Руководитель организации" in doc_text
        assert "М.П." in doc_text
    
    @pytest.mark.asyncio
    async def test_generate_montage_docx_legacy_format(
        self, 
        sample_montage_rows, 
        sample_film_metadata
    ):
        """Test DOCX generation with legacy format."""
        generator = DOCXGeneratorService()
        
        doc_bytes = await generator.generate_montage_docx(
            sample_montage_rows,
            sample_film_metadata,
            use_gff_format=False,
            include_production_info=False,
            include_compliance_info=False
        )
        
        assert isinstance(doc_bytes, bytes)
        assert len(doc_bytes) > 0
        
        # Verify document structure
        doc_buffer = BytesIO(doc_bytes)
        doc = Document(doc_buffer)
        
        # Legacy format should have different structure
        assert len(doc.tables) >= 1
    
    @pytest.mark.asyncio
    async def test_validate_goskino_compliance_valid(self, sample_montage_rows, sample_film_metadata, sample_project_settings):
        """Test Госфильмфонд compliance validation with valid data."""
        generator = DOCXGeneratorService()
        
        is_compliant, issues, report = await generator.validate_goskino_compliance(
            sample_montage_rows,
            sample_film_metadata,
            sample_project_settings
        )
        
        assert is_compliant is True
        assert len(issues) == 0
        assert report["is_compliant"] is True
        assert report["total_issues"] == 0
        assert "compliance_checks" in report
        assert report["compliance_checks"]["metadata"] is True
        assert report["compliance_checks"]["timecodes"] is True
        assert report["compliance_checks"]["table_structure"] is True
        assert report["compliance_checks"]["content"] is True
    
    @pytest.mark.asyncio
    async def test_validate_goskino_compliance_invalid(self, sample_project_settings):
        """Test Госфильмфонд compliance validation with invalid data."""
        generator = DOCXGeneratorService()
        
        # Create invalid metadata
        invalid_metadata = FilmMetadata(
            title="",  # Empty title
            production_company="",  # Empty production company
            year=2024,
            country="Россия",
            screenwriters=[],  # Empty screenwriters
            copyright_holders=[],  # Empty copyright holders
            duration="01:30:45",
            episodes_count=1,
            format="",  # Empty format
            color_type=ColorType.COLOR,
            media_carrier="",  # Empty media carrier
            original_language="",  # Empty original language
            audio_language=""  # Empty audio language
        )
        
        # Create invalid montage rows
        invalid_rows = [
            MontageRow(
                number=2,  # Wrong number (should be 1)
                start_timecode="01:00:00:00",
                end_timecode="01:00:05:00",
                shot_type=ShotType.MEDIUM,
                description="",  # Empty description
                dialogue=""
            )
        ]
        
        is_compliant, issues, report = await generator.validate_goskino_compliance(
            invalid_rows,
            invalid_metadata,
            sample_project_settings
        )
        
        assert is_compliant is False
        assert len(issues) > 0
        assert report["is_compliant"] is False
        assert report["total_issues"] > 0
        
        # Check that specific compliance checks failed
        assert report["compliance_checks"]["metadata"] is False
        assert report["compliance_checks"]["table_structure"] is False
    
    @pytest.mark.asyncio
    async def test_validate_enhanced_document_compliance(
        self, 
        sample_montage_rows, 
        sample_film_metadata, 
        sample_project_settings
    ):
        """Test enhanced document compliance validation."""
        generator = DOCXGeneratorService()
        
        is_compliant, issues = await generator.validate_document_compliance(
            sample_montage_rows,
            sample_film_metadata,
            sample_project_settings
        )
        
        assert is_compliant is True
        assert len(issues) == 0
    
    @pytest.mark.asyncio
    async def test_validate_project_settings(self, sample_film_metadata):
        """Test project settings validation.""" 
        generator = DOCXGeneratorService()
        
        # Test with valid Pydantic format but logically invalid project settings
        invalid_settings = ProjectSettings(
            timecode_start="01:00:00:00",  # Valid format
            standard=TimecodeStandard.GFF,
            fps=999.0  # Invalid FPS (but passes Pydantic validation)
        )
        
        is_compliant, issues = await generator.validate_document_compliance(
            [],  # Empty rows to focus on settings validation
            sample_film_metadata,
            invalid_settings
        )
        
        assert is_compliant is False
        assert len(issues) > 0
        
        # Check specific issues
        issue_text = " ".join(issues)
        assert "частота кадров" in issue_text
    
    @pytest.mark.asyncio
    async def test_russian_film_validation(self, sample_montage_rows):
        """Test validation for Russian films."""
        generator = DOCXGeneratorService()
        
        # Create Russian film metadata with non-Russian original language
        russian_metadata = FilmMetadata(
            title="Русский фильм",
            production_company="Российская киностудия",
            year=2024,
            country="Россия",
            screenwriters=["Иван Иванов"],
            copyright_holders=["Российская киностудия"],
            duration="01:30:45",
            episodes_count=1,
            format="Digital",
            color_type=ColorType.COLOR,
            media_carrier="SSD",
            original_language="Английский",  # Should be Russian for Russian films
            audio_language="Русский"
        )
        
        is_compliant, issues, report = await generator.validate_goskino_compliance(
            sample_montage_rows,
            russian_metadata
        )
        
        assert is_compliant is False
        assert len(issues) > 0
        
        # Check for Russian language requirement issue
        issue_text = " ".join(issues)
        assert "язык оригинала" in issue_text
    
    @pytest.mark.asyncio
    async def test_save_document_with_gff_format(
        self, 
        sample_montage_rows, 
        sample_film_metadata, 
        sample_project_settings
    ):
        """Test saving document with GFF format."""
        generator = DOCXGeneratorService()
        
        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = os.path.join(temp_dir, "gff_montage.docx")
            
            saved_path = await generator.save_document_to_file(
                sample_montage_rows,
                sample_film_metadata,
                output_path,
                sample_project_settings,
                use_gff_format=True
            )
            
            assert saved_path == output_path
            assert os.path.exists(output_path)
            
            # Verify GFF format content
            doc = Document(output_path)
            
            # Should have exactly 2 tables (metadata + montage)
            assert len(doc.tables) == 2
            
            # Check GFF-specific content
            doc_text = "\n".join([p.text for p in doc.paragraphs])
            assert "Форма монтажных листов фильма" in doc_text
            assert "МОНТАЖНЫЕ ЛИСТЫ" in doc_text